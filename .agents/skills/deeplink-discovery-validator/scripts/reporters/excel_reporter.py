"""
Enterprise Multi-Sheet Excel (.xlsx) Report Generator for CarDekho Deep Links.
Builds a professional 18-column inventory report with Executive KPIs, status styling, and evidence links.
"""

import os
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from typing import List, Dict, Any


COLUMNS_SCHEMA = [
    "Deep Link",
    "Normalized Deep Link",
    "Module",
    "Screen",
    "Source",
    "Discovery Method",
    "Expected Screen",
    "Actual Screen",
    "Status",
    "Redirect URL",
    "HTTP Status",
    "App Launch Status",
    "Build Version",
    "Device",
    "OS Version",
    "Last Verified",
    "Evidence Path",
    "Remarks"
]

# Color Palette
COLOR_HEADER_BG = "1F4E79"        # Dark Navy Blue
COLOR_HEADER_FG = "FFFFFF"        # White text
COLOR_ACTIVE_BG = "E2EFDA"        # Soft Green
COLOR_ACTIVE_FG = "375623"
COLOR_CHANGED_BG = "FFF2CC"       # Soft Yellow
COLOR_CHANGED_FG = "7F6000"
COLOR_BROKEN_BG = "FCE4D6"        # Soft Red
COLOR_BROKEN_FG = "C65911"
COLOR_REDIRECT_BG = "DDEBF7"      # Soft Blue
COLOR_REDIRECT_FG = "1F4E78"
COLOR_NEW_BG = "E8D0F5"           # Soft Purple
COLOR_NEW_FG = "5B1E78"
COLOR_DUPLICATE_BG = "EDEDED"     # Soft Gray
COLOR_DUPLICATE_FG = "595959"
COLOR_UNVERIFIED_BG = "F2F2F2"    # Light Gray


class ExcelReporter:
    """Generates standardized, styled Excel workbooks for Deep Link validations."""

    @classmethod
    def generate_report(
        cls,
        records: List[Dict[str, Any]],
        summary_stats: Dict[str, Any],
        output_path: str
    ) -> str:
        os.makedirs(os.path.dirname(os.path.abspath(output_path)), exist_ok=True)
        wb = openpyxl.Workbook()

        # Remove default sheet
        default_sheet = wb.active

        # Sheet 1: Executive Summary Dashboard
        ws_summary = wb.create_sheet(title="Executive_Summary")
        cls._build_summary_sheet(ws_summary, records, summary_stats)

        # Sheet 2: Deep Link Master Inventory
        ws_master = wb.create_sheet(title="Deep_Link_Inventory")
        cls._populate_data_sheet(ws_master, records, title="CarDekho Master Deep Link Inventory")

        # Sheet 3: Diff vs Baseline (NEW, CHANGED, REMOVED, BROKEN)
        diff_records = [r for r in records if r.get("Status") in ("NEW", "CHANGED", "REMOVED", "BROKEN", "REDIRECTED")]
        ws_diff = wb.create_sheet(title="Diff_vs_Baseline")
        cls._populate_data_sheet(ws_diff, diff_records, title="Regression Diff vs Baseline Sheet")

        # Sheet 4: Broken & Exceptions
        broken_records = [r for r in records if r.get("Status") == "BROKEN" or "ERROR" in str(r.get("HTTP Status", ""))]
        ws_broken = wb.create_sheet(title="Broken_And_Exceptions")
        cls._populate_data_sheet(ws_broken, broken_records, title="Broken Deep Links & Failed App Launches")

        # Remove initial blank sheet
        if default_sheet in wb.worksheets and len(wb.worksheets) > 1:
            wb.remove(default_sheet)

        wb.save(output_path)
        return output_path

    @classmethod
    def _build_summary_sheet(cls, ws, records: List[Dict[str, Any]], stats: Dict[str, Any]):
        ws.views.sheetView[0].showGridLines = True

        # Title Banner
        ws.merge_cells("A1:F2")
        title_cell = ws["A1"]
        title_cell.value = "CarDekho Android Deep Link Audit & Validation Dashboard"
        title_cell.font = Font(name="Calibri", size=16, bold=True, color="FFFFFF")
        title_cell.fill = PatternFill(start_color=COLOR_HEADER_BG, end_color=COLOR_HEADER_BG, fill_type="solid")
        title_cell.alignment = Alignment(horizontal="center", vertical="center")

        # KPI Metric Cards
        total_links = len(records)
        active_count = sum(1 for r in records if r.get("Status") == "ACTIVE")
        broken_count = sum(1 for r in records if r.get("Status") == "BROKEN")
        changed_count = sum(1 for r in records if r.get("Status") == "CHANGED")
        new_count = sum(1 for r in records if r.get("Status") == "NEW")
        duplicate_count = sum(1 for r in records if r.get("Status") == "DUPLICATE")

        kpis = [
            ("Total Links", total_links, "4F81BD"),
            ("Active (Passing)", active_count, "375623"),
            ("Broken (Failing)", broken_count, "C65911"),
            ("Changed / Divergent", changed_count, "7F6000"),
            ("Newly Discovered", new_count, "5B1E78"),
            ("Duplicates", duplicate_count, "595959")
        ]

        row_kpi_label = 4
        row_kpi_val = 5

        col_letters = ["A", "B", "C", "D", "E", "F"]
        for idx, (label, val, color) in enumerate(kpis):
            col = col_letters[idx]
            c_lbl = ws[f"{col}{row_kpi_label}"]
            c_lbl.value = label
            c_lbl.font = Font(name="Calibri", size=10, bold=True, color="595959")
            c_lbl.alignment = Alignment(horizontal="center", vertical="center")

            c_val = ws[f"{col}{row_kpi_val}"]
            c_val.value = val
            c_val.font = Font(name="Calibri", size=18, bold=True, color=color)
            c_val.alignment = Alignment(horizontal="center", vertical="center")
            c_val.fill = PatternFill(start_color="F2F2F2", end_color="F2F2F2", fill_type="solid")

        # Breakdown Table
        ws["A7"].value = "Status Breakdown Summary"
        ws["A7"].font = Font(name="Calibri", size=13, bold=True, color="1F4E79")

        headers = ["Status Classification", "Count", "Percentage (%)", "Description / Action Required"]
        for c_idx, h in enumerate(headers, 1):
            cell = ws.cell(row=8, column=c_idx, value=h)
            cell.font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
            cell.fill = PatternFill(start_color="333333", end_color="333333", fill_type="solid")
            cell.alignment = Alignment(horizontal="center" if c_idx in (2, 3) else "left")

        breakdown_data = [
            ("ACTIVE", active_count, "Verified live & opens expected CarDekho screen"),
            ("CHANGED", changed_count, "Opens a different screen or modified route vs baseline"),
            ("BROKEN", broken_count, "App crashed, HTTP 404, or intent failed to resolve"),
            ("NEW", new_count, "Newly discovered deep link route not in historical baseline"),
            ("DUPLICATE", duplicate_count, "Redundant URL alias mapped to existing canonical key"),
            ("NOT_VERIFIABLE", sum(1 for r in records if r.get("Status") == "NOT_VERIFIABLE"), "Static schema unverified due to missing device or dynamic tokens")
        ]

        for r_idx, (st, count, desc) in enumerate(breakdown_data, 9):
            pct = round((count / total_links * 100), 1) if total_links > 0 else 0.0
            ws.cell(row=r_idx, column=1, value=st).font = Font(bold=True)
            ws.cell(row=r_idx, column=2, value=count).alignment = Alignment(horizontal="center")
            ws.cell(row=r_idx, column=3, value=f"{pct}%").alignment = Alignment(horizontal="center")
            ws.cell(row=r_idx, column=4, value=desc)

        # Auto column width
        for col in ws.columns:
            max_len = max(len(str(cell.value or '')) for cell in col)
            col_letter = get_column_letter(col[0].column)
            ws.column_dimensions[col_letter].width = max(max_len + 3, 14)

    @classmethod
    def _populate_data_sheet(cls, ws, records: List[Dict[str, Any]], title: str = "Deep Link Report"):
        ws.views.sheetView[0].showGridLines = True

        # Header Row
        for col_idx, col_name in enumerate(COLUMNS_SCHEMA, 1):
            cell = ws.cell(row=1, column=col_idx, value=col_name)
            cell.font = Font(name="Calibri", size=11, bold=True, color=COLOR_HEADER_FG)
            cell.fill = PatternFill(start_color=COLOR_HEADER_BG, end_color=COLOR_HEADER_BG, fill_type="solid")
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=False)

        # Populate Data Rows
        thin_border = Border(
            left=Side(style='thin', color='E0E0E0'),
            right=Side(style='thin', color='E0E0E0'),
            top=Side(style='thin', color='E0E0E0'),
            bottom=Side(style='thin', color='E0E0E0')
        )

        for row_idx, r in enumerate(records, 2):
            status = r.get("Status", "ACTIVE")
            
            # Select row color styling based on status
            fill_bg = "FFFFFF"
            font_color = "000000"
            if status == "ACTIVE":
                fill_bg = COLOR_ACTIVE_BG
                font_color = COLOR_ACTIVE_FG
            elif status == "CHANGED":
                fill_bg = COLOR_CHANGED_BG
                font_color = COLOR_CHANGED_FG
            elif status == "BROKEN":
                fill_bg = COLOR_BROKEN_BG
                font_color = COLOR_BROKEN_FG
            elif status == "REDIRECTED":
                fill_bg = COLOR_REDIRECT_BG
                font_color = COLOR_REDIRECT_FG
            elif status == "NEW":
                fill_bg = COLOR_NEW_BG
                font_color = COLOR_NEW_FG
            elif status == "DUPLICATE":
                fill_bg = COLOR_DUPLICATE_BG
                font_color = COLOR_DUPLICATE_FG

            for col_idx, col_name in enumerate(COLUMNS_SCHEMA, 1):
                val = r.get(col_name, "")
                cell = ws.cell(row=row_idx, column=col_idx, value=str(val) if val is not None else "")
                cell.border = thin_border
                cell.font = Font(name="Calibri", size=10)

                # Highlight Status column
                if col_name == "Status":
                    cell.fill = PatternFill(start_color=fill_bg, end_color=fill_bg, fill_type="solid")
                    cell.font = Font(name="Calibri", size=10, bold=True, color=font_color)
                    cell.alignment = Alignment(horizontal="center")
                elif col_name in ("HTTP Status", "App Launch Status"):
                    cell.alignment = Alignment(horizontal="center")

        # Auto-fit column widths
        for col in ws.columns:
            max_len = 0
            for cell in col:
                val_str = str(cell.value or '')
                if len(val_str) > max_len:
                    max_len = len(val_str)
            col_letter = get_column_letter(col[0].column)
            ws.column_dimensions[col_letter].width = min(max(max_len + 3, 12), 45)

        # Freeze Header Row
        ws.freeze_panes = "A2"
