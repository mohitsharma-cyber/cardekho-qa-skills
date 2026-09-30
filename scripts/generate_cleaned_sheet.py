import os
import json
import pandas as pd
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

os.makedirs('output/cleaned_google_sheet', exist_ok=True)

# 1. Load original CSV
df_orig = pd.read_csv('input/downloaded_google_sheet.csv', skiprows=1)

# 2. Load audit results JSON
with open('output/google_sheet_audit/deeplink_inventory.json', 'r', encoding='utf-8') as f:
    audit_data = json.load(f)

audit_map = {}
for r in audit_data['inventory']:
    orig_link = r.get('Deep Link', '').strip()
    if orig_link:
        audit_map[orig_link] = r

cleaned_rows = []
for idx, row in df_orig.iterrows():
    raw_url = str(row.get('URLs', '')).strip()
    if not raw_url or raw_url.lower() == 'nan':
        continue
    
    page_type = str(row.get('Type', '')).strip()
    qa_status = str(row.get('QA Status Android', '')).strip()
    view_type = str(row.get('View', '')).strip()
    
    audit_item = audit_map.get(raw_url, {})
    status = audit_item.get('Status', 'ACTIVE')
    http_code = audit_item.get('HTTP Status', '200')
    redirect_url = audit_item.get('Redirect URL', '')
    remarks = audit_item.get('Remarks', '')
    
    # Compute Updated Canonical URL
    updated_url = raw_url
    if status == 'REDIRECTED' and redirect_url:
        if redirect_url.startswith('/'):
            updated_url = f'https://www.cardekho.com{redirect_url}'
        elif redirect_url.startswith('http'):
            updated_url = redirect_url
    elif raw_url == 'www.cardekho.com/ratenow':
        updated_url = 'https://www.cardekho.com/ratenow'
        status = 'FORMAT_FIXED'
        remarks = 'Added https:// protocol prefix'

    final_status = status
    if status == 'ACTIVE':
        final_status = 'WORKING (200 OK)'
    elif status == 'REDIRECTED':
        final_status = 'REDIRECTED (Updated)'
    elif status == 'NOT_VERIFIABLE':
        final_status = 'BROKEN / DEAD'
    elif status == 'DUPLICATE':
        final_status = 'DUPLICATE'

    cleaned_rows.append({
        'S.No': len(cleaned_rows) + 1,
        'Type / Feature': page_type,
        'Original URL': raw_url,
        'Updated / Canonical URL': updated_url,
        'Validation Status': final_status,
        'HTTP Code': http_code,
        'View': view_type,
        'QA Status Android': qa_status,
        'Issue / Audit Remarks': remarks
    })

df_cleaned = pd.DataFrame(cleaned_rows)

# Save CSV for Google Sheets import
csv_out = 'output/cleaned_google_sheet/Updated_CarDekho_DeepLinks.csv'
df_cleaned.to_csv(csv_out, index=False, encoding='utf-8-sig')
print(f'Saved CSV: {csv_out}')

# Build Excel Workbook
xlsx_out = 'output/cleaned_google_sheet/Updated_CarDekho_DeepLinks.xlsx'
wb = openpyxl.Workbook()

# Sheet 1: Master Cleaned Sheet
ws = wb.active
ws.title = 'Updated_Deep_Links'
ws.views.sheetView[0].showGridLines = True

headers = list(df_cleaned.columns)

# Styles
hdr_fill = PatternFill(start_color='1F4E79', end_color='1F4E79', fill_type='solid')
hdr_font = Font(name='Calibri', size=11, bold=True, color='FFFFFF')

thin_border = Border(
    left=Side(style='thin', color='D3D3D3'),
    right=Side(style='thin', color='D3D3D3'),
    top=Side(style='thin', color='D3D3D3'),
    bottom=Side(style='thin', color='D3D3D3')
)

for col_idx, h in enumerate(headers, 1):
    c = ws.cell(row=1, column=col_idx, value=h)
    c.fill = hdr_fill
    c.font = hdr_font
    c.alignment = Alignment(horizontal='center', vertical='center')

for row_idx, r in enumerate(cleaned_rows, 2):
    st = r['Validation Status']
    
    # Status color
    st_fill = PatternFill(start_color='FFFFFF', end_color='FFFFFF', fill_type='solid')
    st_font = Font(name='Calibri', size=10, bold=False)
    
    if 'WORKING' in st:
        st_fill = PatternFill(start_color='E2EFDA', end_color='E2EFDA', fill_type='solid')
        st_font = Font(name='Calibri', size=10, bold=True, color='375623')
    elif 'REDIRECTED' in st:
        st_fill = PatternFill(start_color='DDEBF7', end_color='DDEBF7', fill_type='solid')
        st_font = Font(name='Calibri', size=10, bold=True, color='1F4E78')
    elif 'BROKEN' in st:
        st_fill = PatternFill(start_color='FCE4D6', end_color='FCE4D6', fill_type='solid')
        st_font = Font(name='Calibri', size=10, bold=True, color='C65911')
    elif 'DUPLICATE' in st:
        st_fill = PatternFill(start_color='EDEDED', end_color='EDEDED', fill_type='solid')
        st_font = Font(name='Calibri', size=10, bold=True, color='595959')

    for col_idx, h in enumerate(headers, 1):
        val = r[h]
        cell = ws.cell(row=row_idx, column=col_idx, value=str(val))
        cell.border = thin_border
        cell.font = Font(name='Calibri', size=10)
        
        if h == 'Validation Status':
            cell.fill = st_fill
            cell.font = st_font
            cell.alignment = Alignment(horizontal='center')
        elif h in ('S.No', 'HTTP Code', 'View', 'QA Status Android'):
            cell.alignment = Alignment(horizontal='center')

for col in ws.columns:
    max_len = max(len(str(cell.value or '')) for cell in col)
    col_letter = get_column_letter(col[0].column)
    ws.column_dimensions[col_letter].width = min(max(max_len + 3, 12), 48)

ws.freeze_panes = 'A2'

# Sheet 2: Broken & Action Items
ws_broken = wb.create_sheet(title='Broken_Action_Items')
ws_broken.views.sheetView[0].showGridLines = True
broken_items = [r for r in cleaned_rows if 'BROKEN' in r['Validation Status']]

for col_idx, h in enumerate(headers, 1):
    c = ws_broken.cell(row=1, column=col_idx, value=h)
    c.fill = PatternFill(start_color='C65911', end_color='C65911', fill_type='solid')
    c.font = hdr_font
    c.alignment = Alignment(horizontal='center')

for row_idx, r in enumerate(broken_items, 2):
    for col_idx, h in enumerate(headers, 1):
        c = ws_broken.cell(row=row_idx, column=col_idx, value=str(r[h]))
        c.border = thin_border
        c.font = Font(name='Calibri', size=10)
        if h == 'Validation Status':
            c.fill = PatternFill(start_color='FCE4D6', end_color='FCE4D6', fill_type='solid')
            c.font = Font(name='Calibri', size=10, bold=True, color='C65911')

for col in ws_broken.columns:
    max_len = max(len(str(cell.value or '')) for cell in col)
    col_letter = get_column_letter(col[0].column)
    ws_broken.column_dimensions[col_letter].width = min(max(max_len + 3, 12), 48)

ws_broken.freeze_panes = 'A2'

wb.save(xlsx_out)
print(f'Saved Excel: {xlsx_out}')
print(f'Total Cleaned Rows: {len(cleaned_rows)}')
